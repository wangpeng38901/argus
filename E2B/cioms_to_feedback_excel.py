#!/usr/bin/env python3
"""
将 CIOMS PDF 内容映射到数据反馈 Excel 模板。

功能：
1. 从 URL 或本地路径加载 CIOMS PDF 和 Excel 模板；
2. 判断 PDF 是否为电子档（可提取文本），若疑似扫描件则报错；
3. 解析 CIOMS 常见字段；
4. 通过 JSON/YAML 字段映射配置，将结果写入 Excel 模板并保存输出文件。
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from pypdf import PdfReader

try:
    import requests  # type: ignore
except Exception:
    requests = None


DEFAULT_PDF_URL = ""
DEFAULT_TEMPLATE_URL = "./数据反馈结果模板.xlsx"
DEFAULT_MAPPING_FILE_NAME = "cioms_field_mapping.json"


def get_default_mapping_path() -> Path:
    """
    获取默认映射配置路径：
    - 源码运行：脚本同目录 cioms_field_mapping.json
    - PyInstaller 运行：优先 _MEIPASS 中打包资源，其次 exe 同目录
    """
    if getattr(sys, "frozen", False):
        meipass = getattr(sys, "_MEIPASS", "")
        if meipass:
            bundled = Path(meipass) / DEFAULT_MAPPING_FILE_NAME
            if bundled.exists():
                return bundled
        return Path(sys.executable).resolve().with_name(DEFAULT_MAPPING_FILE_NAME)
    return Path(__file__).resolve().with_name(DEFAULT_MAPPING_FILE_NAME)


@dataclass
class DrugUsage:
    seq: int
    generic_name: str = ""
    trade_name: str = ""
    dosage_form: str = ""
    dose: str = ""
    dose_unit: str = ""
    frequency: str = ""
    days: str = ""
    route: str = ""
    start_date: str = ""
    end_date: str = ""
    indication: str = ""
    kind: str = "怀疑"  # 怀疑 / 合并


@dataclass
class LabCheck:
    check_date: str = ""
    check_name: str = ""
    assessment: str = ""
    result_value: str = ""
    result_unit: str = ""
    normal_low: str = ""
    normal_high: str = ""


@dataclass
class CiomsData:
    report_code: str = ""
    report_type: str = ""
    report_source: str = ""
    organization_type: str = ""
    gender: str = ""
    age: str = ""
    age_unit: str = "岁"
    ethnicity: str = ""
    weight: str = ""
    event_date: str = ""
    reaction_name: str = ""
    reaction_names: List[str] = field(default_factory=list)
    reaction_description: str = ""
    outcome: str = ""
    severe_criteria: List[str] = field(default_factory=list)
    dechallenge: str = ""
    rechallenge: str = ""
    reporter_assessment: str = ""
    company_assessment: str = ""
    report_date: str = ""
    info_source: str = ""
    notes: str = ""
    primary_disease: str = ""
    important_info: str = ""
    lab_items: List[LabCheck] = field(default_factory=list)
    suspected_drugs: List[DrugUsage] = field(default_factory=list)
    concomitant_drugs: List[DrugUsage] = field(default_factory=list)


@dataclass
class ProcessResult:
    page_count: int
    output_path: Path
    report_code: str
    suspected_drug_names: str
    processed_at: str


def download_binary(source: str) -> bytes:
    parts = urllib.parse.urlsplit(source)
    scheme = (parts.scheme or "").lower()

    # 仅 http/https 走网络下载；其余一律按本地路径处理。
    # 这样可避免 Windows 绝对路径 D:\... 被误判为 URL (scheme='d')。
    if scheme in ("http", "https"):
        if requests is not None:
            response = requests.get(source, timeout=30)
            response.raise_for_status()
            return response.content

        safe_path = urllib.parse.quote(parts.path, safe="/")
        encoded_url = urllib.parse.urlunsplit((parts.scheme, parts.netloc, safe_path, parts.query, parts.fragment))
        with urllib.request.urlopen(encoded_url, timeout=30) as resp:
            return resp.read()

    if scheme == "file":
        local_path = Path(urllib.request.url2pathname(parts.path))
    else:
        local_path = Path(source).expanduser()

    if not local_path.is_absolute():
        local_path = (Path.cwd() / local_path).resolve()
    if not local_path.exists():
        raise FileNotFoundError(f"本地文件不存在：{local_path}")
    return local_path.read_bytes()


def normalize_spaces(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text


def to_iso_date(cn_date: str) -> str:
    m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", cn_date)
    if not m:
        return ""
    year, month, day = m.groups()
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def day_span_inclusive(start_iso: str, end_iso: str) -> str:
    try:
        if not start_iso or not end_iso:
            return ""
        s = datetime.strptime(start_iso, "%Y-%m-%d")
        e = datetime.strptime(end_iso, "%Y-%m-%d")
        if e < s:
            return ""
        return str((e - s).days + 1)
    except Exception:
        return ""


def parse_frequency_count_and_cycle_days(freq_raw: str) -> Tuple[str, str]:
    """
    返回 (用药-次数, 周期日数)。
    例：
    - 每3周1次 -> ("1", "21")
    - q3w -> ("1", "21")
    - qd -> ("1", "1")
    - bid -> ("2", "1")
    """
    txt = (freq_raw or "").strip().lower().replace(" ", "")
    if not txt:
        return "", ""

    # 每N周M次 / 每N天M次
    m = re.search(r"每(\d+)周(\d+)次", txt)
    if m:
        weeks = int(m.group(1))
        count = m.group(2)
        return count, str(weeks * 7)
    m = re.search(r"每(\d+)天(\d+)次", txt)
    if m:
        days = int(m.group(1))
        count = m.group(2)
        return count, str(days)

    # q3w / q2d / q8h
    m = re.search(r"q(\d+)([wdh])", txt)
    if m:
        n = int(m.group(1))
        unit = m.group(2)
        if unit == "w":
            return "1", str(n * 7)
        if unit == "d":
            return "1", str(n)
        if unit == "h":
            # 每 n 小时一次，折算日频次
            times_per_day = max(1, round(24 / n))
            return str(times_per_day), "1"

    # qd / bid / tid / qid / once
    if "qid" in txt:
        return "4", "1"
    if "tid" in txt:
        return "3", "1"
    if "bid" in txt:
        return "2", "1"
    if "qd" in txt or "once" in txt:
        return "1", "1"

    # 兜底：抓“X次”
    m = re.search(r"(\d+)\s*次", txt)
    if m:
        return m.group(1), ""

    return "", ""


def parse_cycle_days_from_frequency_text(freq_raw: str) -> str:
    txt = (freq_raw or "").strip().lower().replace(" ", "")
    if not txt:
        return ""
    m = re.search(r"每(\d+)周(\d+)次", txt)
    if m:
        return str(int(m.group(1)) * 7)
    m = re.search(r"每(\d+)天(\d+)次", txt)
    if m:
        return str(int(m.group(1)))
    m = re.search(r"q(\d+)([wdh])", txt)
    if m:
        n = int(m.group(1))
        unit = m.group(2)
        if unit == "w":
            return str(n * 7)
        if unit == "d":
            return str(n)
        if unit == "h":
            return "1"
    return ""


def parse_lab_items_from_text(text: str) -> List[LabCheck]:
    lab_block = extract_first_non_empty_between(
        text,
        [
            ("13. 实验室检查", "13. 相关实验室检查"),
            ("13. 实验室检查", "14-19. 怀疑药物（续）"),
            ("13. 相关实验室检查", "14-19. 怀疑药物（续）"),
        ],
    )
    if not lab_block:
        return []

    # 清理分页噪声：公司编号/附加信息时间戳
    lab_text = re.sub(
        r"公司编号[:：]\s*[A-Z]\d+\s*附加信息\s*\d{4}-\d{2}-\d{2}\s*\d{2}:\d{2}:\d{2}",
        " ",
        lab_block,
    )
    lab_text = re.sub(r"#\s*日期\s*检查/评估/注释\s*结果\s*正常范围\s*高/低", " ", lab_text)
    lab_text = re.sub(r"\s+", " ", lab_text).strip()
    if not lab_text:
        return []

    row_chunks = re.split(r"(?=\b\d{1,2}\s+\d{4}年\d{2}月\d{2}日\b)", lab_text)
    items: List[LabCheck] = []
    for chunk in row_chunks:
        chunk = chunk.strip()
        if not chunk:
            continue

        m = re.match(r"(\d{1,2})\s+(\d{4}年\d{2}月\d{2}日)\s+(.+)$", chunk)
        if not m:
            continue

        date_iso = to_iso_date(m.group(2))
        rest = m.group(3).strip()
        if not rest:
            continue

        # 模式A：名称 [评估] 结果值 单位 正常高 正常低
        ma = re.match(
            r"(.+?)\s+([<>]?\d+(?:\.\d+)?)\s*([^\s]+)\s+([<>]?\d+(?:\.\d+)?)\s+([<>]?\d+(?:\.\d+)?)$",
            rest,
        )
        # 模式B：名称 结果值 单位 正常高 评估 正常低（跨页噪声常见）
        mb = re.match(
            r"(.+?)\s+([<>]?\d+(?:\.\d+)?)\s*([^\s]+)\s+([<>]?\d+(?:\.\d+)?)\s+(升高|降低|正常|异常)\s+([<>]?\d+(?:\.\d+)?)$",
            rest,
        )

        check_name = ""
        assessment = ""
        result_value = ""
        result_unit = ""
        normal_high = ""
        normal_low = ""

        if ma:
            name_and_assess = ma.group(1).strip()
            result_value = ma.group(2).strip()
            result_unit = ma.group(3).strip()
            normal_high = ma.group(4).strip()
            normal_low = ma.group(5).strip()
            m_assess = re.match(r"(.+?)\s+(升高|降低|正常|异常)$", name_and_assess)
            if m_assess:
                check_name = m_assess.group(1).strip()
                assessment = m_assess.group(2).strip()
            else:
                check_name = name_and_assess
        elif mb:
            check_name = mb.group(1).strip()
            result_value = mb.group(2).strip()
            result_unit = mb.group(3).strip()
            normal_high = mb.group(4).strip()
            assessment = mb.group(5).strip()
            normal_low = mb.group(6).strip()
        else:
            continue

        items.append(
            LabCheck(
                check_date=date_iso,
                check_name=check_name,
                assessment=assessment,
                result_value=result_value,
                result_unit=result_unit,
                normal_low=normal_low,
                normal_high=normal_high,
            )
        )

    return items


def extract_between(text: str, start: str, end: str) -> str:
    start_idx = text.find(start)
    if start_idx < 0:
        return ""
    end_idx = text.find(end, start_idx + len(start))
    if end_idx < 0:
        end_idx = len(text)
    return text[start_idx + len(start) : end_idx].strip()


def extract_first_non_empty_between(text: str, candidates: List[Tuple[str, str]]) -> str:
    for start, end in candidates:
        block = extract_between(text, start, end)
        if block.strip():
            return block.strip()
    return ""


def detect_electronic_pdf(pdf_bytes: bytes) -> Tuple[bool, List[str]]:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    page_texts: List[str] = []
    for page in reader.pages:
        txt = (page.extract_text() or "").strip()
        page_texts.append(txt)

    non_empty_pages = sum(1 for t in page_texts if len(t) >= 50)
    total_chars = sum(len(t) for t in page_texts)
    is_electronic = total_chars >= 500 and non_empty_pages >= max(1, len(page_texts) - 1)
    return is_electronic, page_texts


def parse_checkbox_answer(text: str, question: str, candidates: List[str]) -> str:
    idx = text.find(question)
    if idx < 0:
        return ""
    snippet = text[idx : idx + 1200]
    for item in candidates:
        if f"■{item}" in snippet or f"■ {item}" in snippet:
            return item
    return ""


def parse_reaction_names(text: str) -> List[str]:
    block = extract_between(text, "事件报告术语（首位语）（相关症状，如有，用逗号分隔）", "病例描述:")
    if not block:
        return []

    names: List[str] = []
    for raw_line in block.splitlines():
        line = raw_line.strip()
        if not line or "事件报告术语" in line:
            continue

        for part in re.split(r"[，,；;]\s*", line):
            term = part.strip("，,；;。 ")
            if not term:
                continue
            term = re.sub(r"^#?\d+[)）.、]\s*", "", term)
            m = re.search(r"（([^）]+)）", term)
            name = (m.group(1) if m else term).strip()
            if name and name not in names:
                names.append(name)
    return names


def parse_cioms(text: str) -> CiomsData:
    data = CiomsData()
    text = normalize_spaces(text)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    # 公司编号
    m = re.search(r"24b\.\s*生产企业控制编号\.?\s*([A-Z]\d+)", text)
    if not m:
        m = re.search(r"公司编号[:：]\s*([A-Z]\d+)", text)
    if m:
        data.report_code = m.group(1).strip()
        data.notes = data.report_code

    # 报告类型（首次/跟踪）
    m = re.search(r"25a\.\s*报告类型\s*([■□])首次\s*([■□])随访", text)
    if m:
        data.report_type = "首次报告" if m.group(1) == "■" else "跟踪报告"

    # 信息来源
    source_items = []
    for label in ("试验", "文献", "医疗专业", "其他"):
        if re.search(rf"■\s*{label}", text):
            source_items.append(label)
    data.info_source = "；".join(source_items)
    if "医疗专业" in source_items:
        data.organization_type = "医疗机构"

    # 年龄/性别/体重
    m = re.search(r"2a\.\s*年龄\s*(\d+(?:\.\d+)?)\s*岁", text)
    if m:
        data.age = m.group(1)
    m = re.search(r"3\.\s*性别\s*([男女])", text)
    if m:
        data.gender = m.group(1)
    m = re.search(r"3a\.\s*体重\s*(\d+(?:\.\d+)?)\s*公斤", text)
    if m:
        data.weight = m.group(1)

    # 民族/人群（可选）
    m = re.search(r"\d+岁，([^，。\s]{1,8}人)，参加", text)
    if m:
        data.ethnicity = m.group(1).strip()

    # 反应名称（支持多条）
    data.reaction_names = parse_reaction_names(text)
    if data.reaction_names:
        data.reaction_name = data.reaction_names[0]
    else:
        m = re.search(r"事件报告术语.*?\n([^\n]+)", text, re.S)
        if m:
            line = m.group(1).strip()
            name_match = re.search(r"（([^）]+)）", line)
            data.reaction_name = name_match.group(1) if name_match else line

    # 严重性标准
    severe_options = [
        "患者死亡",
        "导致或延长住院时间",
        "永久或重大伤残/丧失工作能力",
        "危及生命",
        "先天畸形",
        "其他",
    ]
    for label in severe_options:
        if re.search(rf"■\s*{re.escape(label)}", text):
            data.severe_criteria.append(label)

    # 事件发生日期（优先正文日期）
    m = re.search(r"(\d{4}年\d{2}月\d{2}日)，可疑感染", text)
    if m:
        data.event_date = to_iso_date(m.group(1))
    if not data.event_date:
        m = re.search(r"年\s*(\d{4})", text)
        if m:
            # 4-6 日期结构被拆分时，尽力拼接
            n = re.search(r"日\s*(\d{1,2}).*?月\s*(\d{1,2}).*?年\s*(\d{4})", text, re.S)
            if n:
                y = int(n.group(3))
                mo = int(n.group(2))
                d = int(n.group(1))
                data.event_date = f"{y:04d}-{mo:02d}-{d:02d}"

    # 病例描述（含续页）
    desc1 = extract_between(text, "病例描述:", " (续附加信息页)")
    desc2 = extract_between(text, "7+13. 不良反应描述（续）", "13. 相关实验室检查")
    data.reaction_description = " ".join(x for x in [desc1, desc2] if x).strip()

    # 转归
    m = re.search(r"结局为[:：]\s*([^。；\n]+)", text)
    if m:
        data.outcome = m.group(1).strip()

    # 停药减轻/再挑战
    data.dechallenge = parse_checkbox_answer(text, "20. 停药后反应减轻了吗？", ["是", "否", "不明", "不适用"])
    data.rechallenge = parse_checkbox_answer(text, "21. 重新用药后是否再次出现", ["是", "否", "不明", "不适用"])

    # 报告日期
    m = re.search(r"本报告日期\s*(\d{4}年\d{2}月\d{2}日)", text)
    if m:
        data.report_date = to_iso_date(m.group(1))

    # 报告人/公司评价
    m = re.search(r"研究者因果关系判断[:：]\s*([\s\S]{0,300})", text)
    if m:
        judge = m.group(1)
        data.reporter_assessment = "可能" if "可能有关" in judge else "待评估"
    m = re.search(r"企业评估[:：]\s*([\s\S]{0,300})", text)
    if m:
        judge = m.group(1)
        data.company_assessment = "可能" if "可能有关" in judge or "预期不良事件" in judge else "待评估"

    # 原患疾病：优先 17. 适应症
    indication_block = extract_between(text, "17. 适应症", "18. 给药日期（从/到）")
    if indication_block:
        m = re.search(r"#\d+\)\s*([^\n\(]+)", indication_block)
        if m:
            data.primary_disease = m.group(1).strip()

    # 重要信息：实验室检查（兼容“13. 实验室检查”与“13. 相关实验室检查”两种标题）
    data.important_info = extract_first_non_empty_between(
        text,
        [
            ("13. 实验室检查", "13. 相关实验室检查"),
            ("13. 实验室检查", "14-19. 怀疑药物（续）"),
            ("13. 相关实验室检查", "14-19. 怀疑药物（续）"),
            ("13. 相关实验室检查", "14. 怀疑药物"),
        ],
    )
    data.lab_items = parse_lab_items_from_text(text)

    # 怀疑用药（主页区块，仅 #1/#2 这组结构化字段最稳定）
    main_suspect_block = extract_between(text, "14. 怀疑药物（包括通用名称）", "20. 停药后反应减轻了吗？")

    drugs: Dict[int, DrugUsage] = {}
    suspect_ids = set()
    for line in [ln.strip() for ln in main_suspect_block.splitlines() if ln.strip()]:
        m = re.match(r"#(\d+)\)\s*([^(#]+?)\(([^)]+)\)\s*(.*)$", line)
        if not m:
            continue
        idx = int(m.group(1))
        suspect_ids.add(idx)
        name = m.group(2).strip()
        trade = m.group(3).strip()
        tail = m.group(4).strip()
        if "给药方案" in tail:
            continue
        item = drugs.get(idx, DrugUsage(seq=idx))
        item.generic_name = item.generic_name or name
        item.trade_name = item.trade_name or trade
        if "注射剂" in tail:
            item.dosage_form = "注射剂"
        elif "冻干粉" in tail:
            item.dosage_form = "冻干粉"
        drugs[idx] = item

    # 剂量（从 15. 日剂量 区块读取）
    dose_block = extract_between(text, "15. 日剂量", "16. 给药途径")
    for m in re.finditer(r"#(\d+)\)\s*([\d.]+)\s*(mg|g|ml|ug|μg|毫克|克|毫升)\s*,\s*([^\n]+)", dose_block):
        idx = int(m.group(1))
        item = drugs.get(idx, DrugUsage(seq=idx))
        item.dose = m.group(2)
        item.dose_unit = m.group(3).replace("μg", "ug")
        freq_raw = m.group(4).strip().rstrip("；;。")
        count_val, cycle_days = parse_frequency_count_and_cycle_days(freq_raw)
        item.frequency = count_val if count_val else freq_raw
        # 周期型频率（如每3周1次）优先映射为给药日数（21）
        if cycle_days:
            item.days = cycle_days
        drugs[idx] = item

    # 给药途径
    route_block = extract_between(text, "16. 给药途径", "17. 适应症")
    for m in re.finditer(r"#(\d+)\)\s*([^\n#]+)", route_block):
        idx = int(m.group(1))
        item = drugs.get(idx, DrugUsage(seq=idx))
        item.route = m.group(2).strip()
        drugs[idx] = item

    # 适应症
    for m in re.finditer(r"#(\d+)\)\s*([^\n#\(]+)\s*(?:\(|$)", indication_block):
        idx = int(m.group(1))
        item = drugs.get(idx, DrugUsage(seq=idx))
        item.indication = m.group(2).strip()
        drugs[idx] = item

    # 开始结束日期（兼容“19. 给药持续时间”与“19. 给药间期”）
    date_block = extract_first_non_empty_between(
        text,
        [
            ("18. 给药日期（从/到）", "19. 给药持续时间"),
            ("18. 给药日期（从/到）", "19. 给药间期"),
        ],
    )
    for m in re.finditer(r"#(\d+)\)\s*(\d{4}年\d{2}月\d{2}日)\s*/\s*(\d{4}年\d{2}月\d{2}日|继续|不明)", date_block):
        idx = int(m.group(1))
        item = drugs.get(idx, DrugUsage(seq=idx))
        item.start_date = to_iso_date(m.group(2))
        end_raw = m.group(3)
        item.end_date = to_iso_date(end_raw) if "年" in end_raw else end_raw
        drugs[idx] = item

    # 持续时间/给药间期（兼容不同模板）
    duration_block = extract_first_non_empty_between(
        text,
        [
            ("19. 给药持续时间", "□是□否□不明"),
            ("19. 给药持续时间", "III.合并药物与病史"),
            ("19. 给药间期", "22. 合并药物和给药日期"),
            ("19. 给药间期", "23. 其他相关病史"),
        ],
    )
    for m in re.finditer(r"#(\d+)\)\s*([\d.]+)\s*day", duration_block):
        idx = int(m.group(1))
        item = drugs.get(idx, DrugUsage(seq=idx))
        # 仅在未有周期天数时，才使用 day 字段兜底
        if not item.days:
            item.days = m.group(2)
        drugs[idx] = item

    # 若未直接解析出用药日数，基于给药起止日期兜底推断（首页数据）
    for idx in list(drugs.keys()):
        item = drugs[idx]
        if (not item.days) and item.start_date and item.end_date and item.end_date not in {"继续", "不明"}:
            item.days = day_span_inclusive(item.start_date, item.end_date)
        drugs[idx] = item

    # 首页怀疑药（每个药物至少输出一条）
    suspected_rows: List[DrugUsage] = []
    for idx in sorted(suspect_ids):
        item = drugs.get(idx)
        if not item or not item.generic_name:
            continue
        item.kind = "怀疑"
        suspected_rows.append(item)

    # 续页 14-19（按“给药方案”逐条展开输出）
    cont_block = extract_first_non_empty_between(
        text,
        [
            ("14-19. 怀疑药物（续）", "22. 合并药物和给药日期（续）"),
            ("14-19. 怀疑药物（续）", "22. 合并药物和给药日期"),
            ("14-19. 怀疑药物（续）", "23. 其他相关病史（续）"),
            ("14-19. 怀疑药物（续）", "23. 其他相关病史"),
            ("14-19. 怀疑药物（续）", "IV.公司信息"),
        ],
    )
    if cont_block:
        cont_lines = [ln.strip() for ln in cont_block.splitlines() if ln.strip()]
        i = 0
        while i < len(cont_lines):
            m = re.match(r"#(\d+)\)\s*([^(#]+?)\(([^)]+)\)\s*(.*)$", cont_lines[i])
            if not m:
                i += 1
                continue
            idx = int(m.group(1))
            suspect_ids.add(idx)
            base = drugs.get(idx, DrugUsage(seq=idx))
            item = DrugUsage(
                seq=idx,
                generic_name=base.generic_name or m.group(2).strip(),
                trade_name=base.trade_name or m.group(3).strip(),
                dosage_form=base.dosage_form,
                route=base.route,
                indication=base.indication,
                kind="怀疑",
            )
            tail = m.group(4).strip()
            if ("注射剂" in tail) and not item.dosage_form:
                item.dosage_form = "注射剂"
            if ("冻干粉" in tail) and not item.dosage_form:
                item.dosage_form = "冻干粉"

            j = i + 1
            while j < len(cont_lines) and not re.match(r"#\d+\)\s*", cont_lines[j]):
                line = cont_lines[j]
                d = re.match(r"([\d.]+)\s*(mg|g|ml|ug|μg|毫克|克|毫升)\s*,\s*([^\n]+)", line)
                if d:
                    item.dose = d.group(1)
                    item.dose_unit = d.group(2).replace("μg", "ug")
                    freq_raw = d.group(3).strip().rstrip("；;。")
                    count_val, cycle_days = parse_frequency_count_and_cycle_days(freq_raw)
                    item.frequency = count_val if count_val else freq_raw
                    if cycle_days:
                        item.days = cycle_days
                elif re.match(r"\d{4}年\d{2}月\d{2}日\s*/\s*(\d{4}年\d{2}月\d{2}日|继续|不明);?$", line):
                    dm = re.match(r"(\d{4}年\d{2}月\d{2}日)\s*/\s*(\d{4}年\d{2}月\d{2}日|继续|不明);?$", line)
                    if dm:
                        item.start_date = to_iso_date(dm.group(1))
                        end_raw = dm.group(2)
                        item.end_date = to_iso_date(end_raw) if "年" in end_raw else end_raw
                elif re.match(r"[\d.]+\s*day", line):
                    if not item.days:
                        item.days = re.match(r"([\d.]+)\s*day", line).group(1)
                elif ("滴注" in line or "口服" in line or "注射" in line) and not item.route:
                    item.route = line
                j += 1

            if (not item.days) and item.start_date and item.end_date and item.end_date not in {"继续", "不明"}:
                item.days = day_span_inclusive(item.start_date, item.end_date)
            if item.generic_name and (item.dose or item.start_date or item.frequency):
                suspected_rows.append(item)
            i = j

    # 去重后输出怀疑药（保留所有给药方案）
    seen_suspected = set()
    for item in suspected_rows:
        key = (
            item.seq,
            item.generic_name,
            item.trade_name,
            item.dosage_form,
            item.dose,
            item.dose_unit,
            item.frequency,
            item.days,
            item.route,
            item.start_date,
            item.end_date,
            item.indication,
        )
        if key in seen_suspected:
            continue
        seen_suspected.add(key)
        data.suspected_drugs.append(item)

    # 合并用药
    in_concomitant = False
    for line in lines:
        if "22. 合并药物和给药日期" in line:
            in_concomitant = True
            continue
        if in_concomitant and (line.startswith("23.") or line.startswith("IV.") or line.startswith("24a.")):
            in_concomitant = False
        if not in_concomitant:
            continue
        m = re.match(r"#(\d+)\)\s*([^(#]+)\(([^)]+)\)\s*;?\s*([^/]+)\s*/\s*(.+)$", line)
        if not m:
            continue
        idx = int(m.group(1))
        item = DrugUsage(
            seq=idx,
            generic_name=m.group(2).strip(),
            trade_name=m.group(3).strip(),
            start_date=to_iso_date(m.group(4).strip()) if "年" in m.group(4) else m.group(4).strip(),
            end_date=to_iso_date(m.group(5).strip()) if "年" in m.group(5) else m.group(5).strip(),
            kind="合并",
        )
        data.concomitant_drugs.append(item)

    if not data.outcome:
        data.outcome = "未好转" if "症状持续" in text else ""

    if not data.organization_type:
        data.organization_type = "医疗机构" if "医疗专业" in data.info_source else ""

    return data


def clear_sheet_from_row(sheet, row_start: int = 2) -> None:
    if sheet.max_row >= row_start:
        sheet.delete_rows(row_start, sheet.max_row - row_start + 1)


def load_mapping_config(config_path: Path) -> Dict[str, Any]:
    suffix = config_path.suffix.lower()
    raw = config_path.read_text(encoding="utf-8")
    if suffix == ".json":
        return json.loads(raw)
    if suffix in {".yaml", ".yml"}:
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise RuntimeError("读取 YAML 映射配置需要安装 pyyaml：pip3 install pyyaml") from exc
        cfg = yaml.safe_load(raw)
        return cfg if isinstance(cfg, dict) else {}
    raise RuntimeError(f"不支持的映射配置格式：{config_path.suffix}（仅支持 .json/.yaml/.yml）")


def cioms_to_context(cioms: CiomsData) -> Dict[str, Any]:
    all_drugs = sorted(cioms.suspected_drugs, key=lambda x: x.seq) + sorted(
        cioms.concomitant_drugs, key=lambda x: x.seq
    )
    cioms_dict = {
        "report_code": cioms.report_code,
        "report_type": cioms.report_type or "首次报告",
        "report_source": cioms.report_source,
        "organization_type": cioms.organization_type,
        "gender": cioms.gender,
        "age": cioms.age,
        "age_unit": cioms.age_unit,
        "ethnicity": cioms.ethnicity,
        "weight": cioms.weight,
        "event_date": cioms.event_date,
        "reaction_name": cioms.reaction_name,
        "reaction_names": cioms.reaction_names,
        "reaction_description": cioms.reaction_description,
        "outcome": cioms.outcome,
        "severe_criteria": cioms.severe_criteria,
        "dechallenge": cioms.dechallenge,
        "rechallenge": cioms.rechallenge,
        "reporter_assessment": cioms.reporter_assessment,
        "company_assessment": cioms.company_assessment,
        "report_date": cioms.report_date,
        "info_source": cioms.info_source,
        "notes": cioms.notes,
        "primary_disease": cioms.primary_disease,
        "important_info": cioms.important_info,
        "lab_items": [x.__dict__ for x in cioms.lab_items],
        "suspected_drugs": [d.__dict__ for d in cioms.suspected_drugs],
        "concomitant_drugs": [d.__dict__ for d in cioms.concomitant_drugs],
    }
    computed = {
        "report_level": "严重" if cioms.severe_criteria else "一般",
        "severe_criteria_joined": "；".join(cioms.severe_criteria),
        "drug_rows": [d.__dict__ for d in all_drugs],
        "reaction_rows": [
            {"reaction_name": name, "reaction_level": "严重" if cioms.severe_criteria else "一般"}
            for name in (cioms.reaction_names if cioms.reaction_names else ([cioms.reaction_name] if cioms.reaction_name else []))
        ],
        "lab_rows": [x.__dict__ for x in cioms.lab_items],
    }

    # 兼容两类配置：
    # 1) 新配置：source 使用 report_type / all_drugs 等扁平 key
    # 2) 旧配置：source 使用 cioms.xxx / computed.xxx / template_feedback_code
    return {
        "cioms": cioms_dict,
        "computed": computed,
        "template_feedback_code": "",
        "feedback_code": "",
        "report_type": cioms_dict["report_type"],
        "report_severity": computed["report_level"],
        "severe_criteria_joined": computed["severe_criteria_joined"],
        "organization_type": cioms.organization_type,
        "gender": cioms.gender,
        "age": cioms.age,
        "age_unit": cioms.age_unit,
        "ethnicity": cioms.ethnicity,
        "weight": cioms.weight,
        "history_adr": "不详",
        "family_adr": "不详",
        "event_date": cioms.event_date,
        "reaction_description": cioms.reaction_description,
        "outcome": cioms.outcome,
        "dechallenge": cioms.dechallenge,
        "rechallenge": cioms.rechallenge,
        "reporter_assessment": cioms.reporter_assessment,
        "company_assessment": cioms.company_assessment,
        "report_date": cioms.report_date,
        "info_source": cioms.info_source,
        "notes": cioms.notes,
        "primary_disease": cioms.primary_disease,
        "reaction_name": cioms.reaction_name,
        "reaction_level": computed["report_level"],
        "important_info": cioms.important_info,
        "all_drugs": computed["drug_rows"],
        "lab_rows": computed["lab_rows"],
    }


def get_by_path(data: Dict[str, Any], path: str) -> Any:
    cur: Any = data
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return None
    return cur


def resolve_value(source: str, row_ctx: Dict[str, Any], global_ctx: Dict[str, Any], default: Any = "") -> Any:
    if source == "":
        return default
    if source.startswith("row."):
        v = get_by_path(row_ctx, source[4:])
        return default if v is None else v
    if source.startswith("global."):
        v = get_by_path(global_ctx, source[7:])
        return default if v is None else v
    if source in row_ctx:
        v = row_ctx.get(source)
        return default if v is None else v
    v = get_by_path(global_ctx, source)
    if v is not None:
        return v
    if source in global_ctx:
        v = global_ctx.get(source)
        return default if v is None else v
    return default


def header_to_col_map(sheet, header_row: int) -> Dict[str, int]:
    return {
        str(sheet.cell(header_row, col).value).strip(): col
        for col in range(1, sheet.max_column + 1)
        if sheet.cell(header_row, col).value is not None and str(sheet.cell(header_row, col).value).strip() != ""
    }


def apply_single_sheet_mapping(sheet, sheet_cfg: Dict[str, Any], global_ctx: Dict[str, Any]) -> None:
    clear_from = int(sheet_cfg.get("clear_from_row", 0) or 0)
    if clear_from > 0:
        clear_sheet_from_row(sheet, clear_from)

    target_row = int(sheet_cfg.get("target_row", 2))
    header_row = int(sheet_cfg.get("header_row", 1))
    header_map = header_to_col_map(sheet, header_row)

    for item in sheet_cfg.get("mappings", []):
        source = item.get("source", "")
        default = item.get("default", "")
        if source == "literal":
            value = item.get("value", default)
        else:
            value = resolve_value(source, {}, global_ctx, default=default)
        col = item.get("column")
        if col is None and item.get("header"):
            col = header_map.get(str(item["header"]))
        if col is None:
            continue
        sheet.cell(target_row, int(col), value=value)


def apply_list_sheet_mapping(sheet, sheet_cfg: Dict[str, Any], global_ctx: Dict[str, Any]) -> None:
    clear_from = int(sheet_cfg.get("clear_from_row", 0) or 0)
    if clear_from > 0:
        clear_sheet_from_row(sheet, clear_from)

    source_key = sheet_cfg.get("source", "")
    rows = resolve_value(source_key, {}, global_ctx, default=[])
    if not isinstance(rows, list):
        return

    start_row = int(sheet_cfg.get("start_row", 2))
    header_row = int(sheet_cfg.get("header_row", 1))
    header_map = header_to_col_map(sheet, header_row)

    for i, row_data in enumerate(rows):
        if not isinstance(row_data, dict):
            continue
        row_idx = start_row + i
        for item in sheet_cfg.get("mappings", []):
            source = item.get("source", "")
            default = item.get("default", "")
            if source == "literal":
                value = item.get("value", default)
            else:
                value = resolve_value(source, row_data, global_ctx, default=default)
            col = item.get("column")
            if col is None and item.get("header"):
                col = header_map.get(str(item["header"]))
            if col is None:
                continue
            sheet.cell(row_idx, int(col), value=value)


def fill_feedback_code_from_template(wb, global_ctx: Dict[str, Any]) -> None:
    # 反馈码填写 24b. 生产企业控制编号（cioms.report_code）
    code = str(get_by_path(global_ctx, "cioms.report_code") or "").strip()
    if not code:
        # 兜底：若 24b 未解析到，则使用模板中的原反馈码
        if "药品不良反应报告表" in wb.sheetnames:
            sheet = wb["药品不良反应报告表"]
            headers = header_to_col_map(sheet, 1)
            col = headers.get("反馈码")
            if col:
                val = sheet.cell(2, col).value
                if val not in (None, ""):
                    code = str(val).strip()
    global_ctx["feedback_code"] = code
    global_ctx["template_feedback_code"] = code


def apply_legacy_rows_mapping(sheet, sheet_cfg: Dict[str, Any], global_ctx: Dict[str, Any]) -> None:
    clear_from = int(sheet_cfg.get("clear_from_row", 0) or 0)
    if clear_from > 0:
        clear_sheet_from_row(sheet, clear_from)

    rows_cfg = sheet_cfg.get("rows", [])
    header_map = header_to_col_map(sheet, 1)
    for row_cfg in rows_cfg:
        target_row = int(row_cfg.get("row", 2))
        by_header = row_cfg.get("by_header", {})
        by_column = row_cfg.get("by_column", {})

        for header, rule in by_header.items():
            if not isinstance(rule, dict):
                continue
            source = str(rule.get("source", ""))
            default = rule.get("default", "")
            value = rule.get("value", default) if source == "literal" else resolve_value(source, {}, global_ctx, default)
            col = header_map.get(str(header))
            if col:
                sheet.cell(target_row, col, value=value)

        for col_text, rule in by_column.items():
            if not isinstance(rule, dict):
                continue
            source = str(rule.get("source", ""))
            default = rule.get("default", "")
            value = rule.get("value", default) if source == "literal" else resolve_value(source, {}, global_ctx, default)
            sheet.cell(target_row, int(col_text), value=value)


def apply_legacy_table_mapping(sheet, sheet_cfg: Dict[str, Any], global_ctx: Dict[str, Any]) -> None:
    clear_from = int(sheet_cfg.get("clear_from_row", 0) or 0)
    if clear_from > 0:
        clear_sheet_from_row(sheet, clear_from)

    table_cfg = sheet_cfg.get("table", {})
    if not isinstance(table_cfg, dict):
        return
    rows = resolve_value(str(table_cfg.get("source", "")), {}, global_ctx, default=[])
    if not isinstance(rows, list):
        return

    start_row = int(table_cfg.get("start_row", 2))
    header_map = header_to_col_map(sheet, 1)
    by_header = table_cfg.get("by_header", {})
    by_column = table_cfg.get("by_column", {})

    for i, row_data in enumerate(rows):
        if not isinstance(row_data, dict):
            continue
        row_idx = start_row + i
        for header, rule in by_header.items():
            if not isinstance(rule, dict):
                continue
            source = str(rule.get("source", ""))
            default = rule.get("default", "")
            value = rule.get("value", default) if source == "literal" else resolve_value(source, row_data, global_ctx, default)
            col = header_map.get(str(header))
            if col:
                sheet.cell(row_idx, col, value=value)

        for col_text, rule in by_column.items():
            if not isinstance(rule, dict):
                continue
            source = str(rule.get("source", ""))
            default = rule.get("default", "")
            value = rule.get("value", default) if source == "literal" else resolve_value(source, row_data, global_ctx, default)
            sheet.cell(row_idx, int(col_text), value=value)


def ensure_sheet_and_headers(wb, sheet_cfg: Dict[str, Any]):
    sheet_name = sheet_cfg.get("name")
    if not sheet_name:
        return None

    sheet = wb[sheet_name] if sheet_name in wb.sheetnames else None
    if sheet is None and sheet_cfg.get("create_if_missing"):
        sheet = wb.create_sheet(title=sheet_name)

    if sheet is None:
        return None

    headers = sheet_cfg.get("headers", [])
    if isinstance(headers, list) and headers:
        for idx, h in enumerate(headers, start=1):
            if h is not None and str(h).strip() != "":
                sheet.cell(1, idx, value=str(h))
    return sheet


def generate_cover_pdf(output_path: Path, rows: List[Tuple[str, str, str]]) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("生成 cover PDF 需要安装 reportlab：pip install reportlab") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    data: List[List[str]] = [["报告编号", "怀疑用药", "本次处理的日期时间"]]
    for report_code, suspected_drugs, processed_at in rows:
        data.append([report_code or "", suspected_drugs or "", processed_at or ""])

    doc = SimpleDocTemplate(str(output_path), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    table = Table(data, colWidths=[170, 180, 170], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
            ]
        )
    )
    doc.build([table])


def write_excel(cioms: CiomsData, template_bytes: bytes, output_path: Path, mapping_cfg: Dict[str, Any]) -> None:
    wb = load_workbook(io.BytesIO(template_bytes))
    context = cioms_to_context(cioms)
    fill_feedback_code_from_template(wb, context)

    sheets_cfg = mapping_cfg.get("sheets", [])
    for sheet_cfg in sheets_cfg:
        sheet = ensure_sheet_and_headers(wb, sheet_cfg)
        if sheet is None:
            continue
        # 兼容旧配置结构（rows/table）与新配置结构（mode/mappings）
        if "rows" in sheet_cfg:
            apply_legacy_rows_mapping(sheet, sheet_cfg, context)
            continue
        if "table" in sheet_cfg:
            apply_legacy_table_mapping(sheet, sheet_cfg, context)
            continue
        mode = sheet_cfg.get("mode", "single")
        if mode == "list":
            apply_list_sheet_mapping(sheet, sheet_cfg, context)
        else:
            apply_single_sheet_mapping(sheet, sheet_cfg, context)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def process_single_pdf(
    pdf_source: str,
    template_source: str,
    mapping_cfg: Dict[str, Any],
    output_path: Path,
) -> ProcessResult:
    pdf_bytes = download_binary(pdf_source)
    ok, page_texts = detect_electronic_pdf(pdf_bytes)
    if not ok:
        raise RuntimeError(f"检测失败：{pdf_source} 疑似扫描件（可提取文本不足），请提供电子档 PDF。")

    full_text = "\n".join(page_texts)
    cioms = parse_cioms(full_text)
    template_bytes = download_binary(template_source)
    write_excel(cioms, template_bytes, output_path, mapping_cfg)
    suspected_name_list: List[str] = []
    for d in cioms.suspected_drugs:
        name = (d.generic_name or d.trade_name or "").strip()
        if name and name not in suspected_name_list:
            suspected_name_list.append(name)
    suspected_drug_names = "；".join(suspected_name_list)

    return ProcessResult(
        page_count=len(page_texts),
        output_path=output_path,
        report_code=cioms.report_code,
        suspected_drug_names=suspected_drug_names,
        processed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="CIOMS PDF -> 数据反馈Excel 映射工具")
    parser.add_argument(
        "--pdf-url",
        "--pdf-source",
        default=DEFAULT_PDF_URL,
        help="CIOMS PDF 来源（支持 URL 或本地路径）。为空时自动处理当前目录全部 PDF",
    )
    parser.add_argument(
        "--template-url",
        "--template-source",
        default=DEFAULT_TEMPLATE_URL,
        help="Excel 模板来源（支持 URL 或本地路径）",
    )
    parser.add_argument(
        "--mapping-config",
        default=str(Path(__file__).resolve().with_name(DEFAULT_MAPPING_FILE_NAME)),
        help="字段映射配置文件路径（支持 .json/.yaml/.yml）",
    )
    parser.add_argument(
        "--output",
        default="output_数据反馈结果-14910028650182082562330.xlsx",
        help="输出 Excel 路径（单文件模式使用）",
    )
    args = parser.parse_args()

    mapping_cfg = load_mapping_config(Path(args.mapping_config).resolve())
    mapping_path = Path(args.mapping_config).resolve()

    # 单文件模式
    if str(args.pdf_url).strip():
        output = Path(args.output).resolve()
        result = process_single_pdf(args.pdf_url, args.template_url, mapping_cfg, output)
        print("处理完成：")
        print(f"- 电子档校验：通过（共 {result.page_count} 页）")
        print(f"- 映射配置：{mapping_path}")
        print(f"- 输出文件：{result.output_path}")
        return

    # 批处理模式：当前目录全部 PDF
    cwd = Path.cwd()
    pdf_files = sorted(cwd.glob("*.pdf"))
    if not pdf_files:
        raise RuntimeError(f"当前目录未找到 PDF 文件：{cwd}")

    print("批处理模式：")
    print(f"- 目录：{cwd}")
    print(f"- 模板：{args.template_url}")
    print(f"- 映射配置：{mapping_path}")
    print(f"- 待处理 PDF 数量：{len(pdf_files)}")

    cover_rows: List[Tuple[str, str, str]] = []
    for pdf in pdf_files:
        out = pdf.with_suffix(".xlsx")
        result = process_single_pdf(str(pdf), args.template_url, mapping_cfg, out)
        cover_rows.append((result.report_code, result.suspected_drug_names, result.processed_at))
        print(f"  · {pdf.name} -> {result.output_path.name} （{result.page_count} 页，电子档校验通过）")

    cover_pdf = cwd / "cover.pdf"
    generate_cover_pdf(cover_pdf, cover_rows)
    print(f"- 汇总封面：{cover_pdf.name}")


if __name__ == "__main__":
    main()
